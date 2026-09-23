"""Fact-check parsing, Word export and CV phone extraction."""

from docx import Document

from src.cv_extractor import CVExtractor
from src.documents import save_document
from src.hybrid_llm import HybridLLM, LLMConfig
from src.writer import GeneratedDocument


def test_fact_check_parses_model_json(monkeypatch):
    llm = HybridLLM(LLMConfig(groq_api_key="x"))
    monkeypatch.setattr(llm, "_call_groq", lambda *a, **k: 'Sure: {"problems": [{"claim": "master\'s thesis", "issue": "it was the undergraduate thesis"}]}')
    problems = llm.fact_check("My master's thesis...", "B.Sc. thesis: chest disease detection")
    assert problems == [{"claim": "master's thesis", "issue": "it was the undergraduate thesis"}]


def test_fact_check_tolerates_bad_output(monkeypatch):
    llm = HybridLLM(LLMConfig(groq_api_key="x"))
    monkeypatch.setattr(llm, "_call_groq", lambda *a, **k: "no json here")
    assert llm.fact_check("text", "sources") == []


def test_hybrid_generate_fixes_flagged_facts(monkeypatch):
    llm = HybridLLM(LLMConfig(groq_api_key="x"))
    monkeypatch.setattr(llm, "_check_local_available", lambda: False)
    calls = []

    def fake_groq(prompt, system="", max_tokens=2048, reasoning="low"):
        calls.append(prompt)
        if prompt.startswith("You are fact-checking"):
            if "master's thesis" in prompt.split("DOCUMENT:")[1]:
                return '{"problems": [{"claim": "master\'s thesis", "issue": "undergraduate thesis"}]}'
            return '{"problems": []}'
        if prompt.startswith("Correct these factual problems"):
            return "My undergraduate thesis used deep learning to detect chest disease in X-ray images."
        return "My master's thesis used deep learning to detect chest disease in X-ray images."

    monkeypatch.setattr(llm, "_call_groq", fake_groq)
    result = llm.hybrid_generate("Write an SOP", "B.Sc. thesis: chest disease", kind="sop")
    assert "undergraduate thesis" in result["humanized"]
    assert result["fact_fixes"] and result["fact_problems"] == []


def test_docx_export_builds_headings_lists_and_tables(tmp_path):
    text = ("# Graph learning for X-rays\n\n## Work plan\n\n- Year 1: baselines\n\n"
            "| Year | Milestone |\n|------|-----------|\n| 1 | Baselines |\n| 2 | Graph models |\n")
    doc = GeneratedDocument("proposal", "Research Proposal", text, post_title="PhD in ML")
    paths = save_document(doc, author="Test User", output_dir=tmp_path)
    word = Document(paths["docx"])
    assert [p.text for p in word.paragraphs if p.style.name.startswith("Heading")] == \
        ["Graph learning for X-rays", "Work plan"]
    assert len(word.tables) == 1 and word.tables[0].cell(2, 1).text == "Graph models"
    assert paths["md"].read_text().startswith("# Graph learning")


def test_phone_ignores_address_numbers():
    cv = "RATUL SAHA\n240/12 President Road, Dhaka 1400\n+880 1910 076810 | me@example.com\n2020 - 2024"
    assert CVExtractor()._extract_phone(cv) == "+880 1910 076810"


def test_fact_check_drops_claims_not_in_document(monkeypatch):
    llm = HybridLLM(LLMConfig(groq_api_key="x"))
    monkeypatch.setattr(llm, "_call_groq", lambda *a, **k: '{"problems": [{"claim": "CGPA: 4.00 / 4.00", "issue": "x"}, {"claim": "first author of LW-ALDDNet", "issue": "Prito is first author"}]}')
    problems = llm.fact_check("I was first author of LW-ALDDNet.", "sources")
    assert [p["claim"] for p in problems] == ["first author of LW-ALDDNet"]


def test_context_includes_whole_cv():
    from src.resource_loader import UserResources
    r = UserResources(name="A", cv_text="x" * 5000 + " MICCAI benchmark")
    assert "MICCAI benchmark" in r.to_context_string()


class _FakeCompletions:
    def __init__(self, limited):
        self.limited, self.models = limited, []

    def create(self, model, **kwargs):
        import httpx
        from groq import RateLimitError
        from types import SimpleNamespace as NS
        self.models.append(model)
        if model in self.limited:
            req = httpx.Request("POST", "https://api.groq.com")
            raise RateLimitError("Rate limit reached ... Please try again in 19m59.2s. org_123",
                                 response=httpx.Response(429, request=req), body=None)
        return NS(choices=[NS(message=NS(content=f"text from {model}"), finish_reason="stop")])


def _llm_with(limited):
    from types import SimpleNamespace as NS
    llm = HybridLLM(LLMConfig(groq_api_key="x", groq_model="big", groq_fallback_model="small"))
    completions = _FakeCompletions(limited)
    llm._groq_client = NS(chat=NS(completions=completions))
    return llm, completions


def test_groq_falls_back_to_second_model_on_rate_limit():
    llm, completions = _llm_with({"big"})
    assert llm._call_groq("hi") == "text from small"
    assert completions.models == ["big", "small"]


def test_groq_rate_limit_message_is_readable_and_hides_org_id():
    import pytest
    llm, _ = _llm_with({"big", "small"})
    with pytest.raises(Exception) as e:
        llm._call_groq("hi")
    assert "usage limit" in str(e.value) and "19m59s" in str(e.value)
    assert "org_" not in str(e.value)
