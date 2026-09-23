"""Save generated SOPs and proposals as Markdown and Word files."""

import re
from datetime import datetime
from pathlib import Path
from typing import Dict

try:
    from docx import Document
    from docx.shared import Pt
except ImportError:
    Document = None

OUTPUT_DIR = Path(__file__).parent.parent / "resources" / "saved_responses"


def _slug(text: str, limit: int = 40) -> str:
    return re.sub(r"[^a-z0-9]+", "-", text.lower()).strip("-")[:limit].strip("-") or "general"


def _add_runs(paragraph, text: str):
    """Add text, turning *italic* / **bold** markdown into Word formatting."""
    for part in re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*)", text):
        if part.startswith("**") and part.endswith("**"):
            paragraph.add_run(part[2:-2]).bold = True
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            paragraph.add_run(part[1:-1]).italic = True
        elif part:
            paragraph.add_run(part)


def to_docx(text: str, title: str, author: str, path: Path):
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)

    lines = text.splitlines()
    if not (lines and lines[0].startswith("# ")):
        doc.add_heading(title, level=1)
    if author:
        doc.add_paragraph(author)

    paragraph_lines = []

    def flush():
        if paragraph_lines:
            _add_runs(doc.add_paragraph(), " ".join(paragraph_lines))
            paragraph_lines.clear()

    table_rows = []

    def flush_table():
        rows = [r for r in table_rows if not re.fullmatch(r"\|?[\s:|-]+\|?", r)]  # drop |---| separators
        table_rows.clear()
        if not rows:
            return
        cells = [[c.strip() for c in r.strip().strip("|").split("|")] for r in rows]
        width = max(len(r) for r in cells)
        table = doc.add_table(rows=len(cells), cols=width)
        table.style = "Table Grid"
        for i, row in enumerate(cells):
            for j, value in enumerate(row):
                table.cell(i, j).text = value
                if i == 0:
                    for run in table.cell(i, j).paragraphs[0].runs:
                        run.bold = True

    for line in lines:
        stripped = line.strip()
        if stripped.startswith("|"):
            flush()
            table_rows.append(stripped)
            continue
        flush_table()
        heading = re.match(r"^(#{1,6})\s+(.*)", stripped)
        bullet = re.match(r"^(?:[-*]|\d+[.)])\s+(.*)", stripped)
        if not stripped:
            flush()
        elif heading:
            flush()
            doc.add_heading(heading.group(2), level=min(len(heading.group(1)), 3))
        elif bullet:
            flush()
            style = "List Number" if stripped[0].isdigit() else "List Bullet"
            _add_runs(doc.add_paragraph(style=style), bullet.group(1))
        else:
            paragraph_lines.append(stripped)
    flush()
    flush_table()
    doc.save(path)


def save_document(doc, author: str = "", output_dir: Path = None) -> Dict[str, Path]:
    """Write doc.text to .md and .docx. Returns {"md": path, "docx": path (if python-docx is installed)}."""
    output_dir = output_dir or OUTPUT_DIR
    output_dir.mkdir(parents=True, exist_ok=True)
    stamp = datetime.now().strftime("%Y%m%d-%H%M")
    base = f"{doc.kind}_{_slug(doc.post_title)}_{stamp}"

    paths = {"md": output_dir / f"{base}.md"}
    paths["md"].write_text(doc.text + "\n", encoding="utf-8")
    if Document is not None:
        paths["docx"] = output_dir / f"{base}.docx"
        to_docx(doc.text, doc.title, author, paths["docx"])
    return paths
